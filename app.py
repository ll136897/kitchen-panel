"""烤肉店后厨备餐系统 - Flask 主应用"""
import math
from flask import Flask, request, jsonify, render_template, g
from models import get_db, init_db
from parser import parse_order_text, match_packages_in_db
from calculator import calc_order_requirements, calc_dashboard, preview_parse, calc_merged_prep, calc_prep_urgency

app = Flask(__name__, template_folder="templates", static_folder="static")

# 模块加载时初始化数据库（保证 gunicorn 多 worker 也能跑）
init_db()

# 尝试从 GitHub 恢复数据库（防止重新部署丢数据）
try:
    from db_backup import restore_from_github, backup_async, start_auto_backup
    restored = restore_from_github()
    if restored:
        print("[init] 已从 GitHub 恢复数据库")
except Exception as _e:
    print(f"[init] restore_from_github skipped: {_e}")

try:
    from seed import seed_data
    seed_data()
except Exception as _e:
    print(f"[init] seed_data skipped: {_e}")

# 启动定时备份（10分钟）
try:
    start_auto_backup(600)
except Exception as _e:
    print(f"[init] auto_backup skipped: {_e}")

# 迁移：去掉"刷子"工具（用户要求菜单无刷子，兼容已有数据库）
try:
    _db = get_db()
    _cur = _db.cursor()
    _cur.execute("SELECT id FROM tools WHERE name = '刷子'")
    _row = _cur.fetchone()
    if _row:
        _bid = _row["id"]
        _cur.execute("DELETE FROM package_tools WHERE tool_id = ?", (_bid,))
        _cur.execute("DELETE FROM tool_loans WHERE tool_id = ?", (_bid,))
        _cur.execute("DELETE FROM tools WHERE id = ?", (_bid,))
        _db.commit()
        print(f"[migrate] 已删除刷子工具(id={_bid})及关联记录")
    _db.close()
except Exception as _e:
    print(f"[migrate] 删除刷子跳过: {_e}")

# 迁移：拆分"卡式炉/烤盘"为"卡式炉"和"烤盘"两个独立工具（兼容已有数据库）
try:
    _db = get_db()
    _cur = _db.cursor()
    _cur.execute("SELECT id FROM tools WHERE name = '卡式炉/烤盘'")
    _row = _cur.fetchone()
    if _row:
        _old_id = _row["id"]
        # 1. 新增两个新工具（卡式炉+烤盘），库存/阈值/成本按合理拆分
        _cur.execute("SELECT stock, threshold, cost FROM tools WHERE id = ?", (_old_id,))
        _old = _cur.fetchone()
        _old_stock = _old["stock"] if _old else 15
        _old_thr = _old["threshold"] if _old else 3
        _cur.execute("INSERT INTO tools (name, stock, threshold, cost) VALUES (?,?,?,?)",
                     ("卡式炉", _old_stock, _old_thr, 65.0))
        _cur.execute("INSERT INTO tools (name, stock, threshold, cost) VALUES (?,?,?,?)",
                     ("烤盘", _old_stock, _old_thr, 15.0))
        _new_stove = _cur.execute("SELECT id FROM tools WHERE name='卡式炉'").fetchone()["id"]
        _new_pan = _cur.execute("SELECT id FROM tools WHERE name='烤盘'").fetchone()["id"]
        # 2. 给每个套餐添加两条 package_tools 关联，per_package 等于原值
        _cur.execute("SELECT package_id, per_package FROM package_tools WHERE tool_id = ?", (_old_id,))
        for _r in _cur.fetchall():
            _cur.execute("INSERT INTO package_tools (package_id, tool_id, per_package) VALUES (?,?,?)",
                         (_r["package_id"], _new_stove, _r["per_package"]))
            _cur.execute("INSERT INTO package_tools (package_id, tool_id, per_package) VALUES (?,?,?)",
                         (_r["package_id"], _new_pan, _r["per_package"]))
        # 3. 删除旧"卡式炉/烤盘"的 package_tools 关联和工具记录
        _cur.execute("DELETE FROM package_tools WHERE tool_id = ?", (_old_id,))
        _cur.execute("DELETE FROM tool_loans WHERE tool_id = ?", (_old_id,))
        _cur.execute("DELETE FROM tools WHERE id = ?", (_old_id,))
        _db.commit()
        print(f"[migrate] 已拆分'卡式炉/烤盘'(id={_old_id})为卡式炉(id={_new_stove})+烤盘(id={_new_pan})")
    _db.close()
except Exception as _e:
    print(f"[migrate] 拆分卡式炉/烤盘跳过: {_e}")

# 迁移：重命名食材"应季水果三样" → "应季水果"（用户指定）
try:
    _db = get_db()
    _cur = _db.cursor()
    _cur.execute("SELECT id FROM ingredients WHERE name = '应季水果三样'")
    _row = _cur.fetchone()
    if _row:
        _cur.execute("UPDATE ingredients SET name = '应季水果' WHERE id = ?", (_row["id"],))
        _db.commit()
        print(f"[migrate] 已重命名食材'应季水果三样'(id={_row['id']})为'应季水果'")
    _db.close()
except Exception as _e:
    print(f"[migrate] 重命名应季水果三样跳过: {_e}")

# 迁移：修复历史订单 booking_date/booking_time 解析错误（旧 parser 对"9月11日12点"等格式解析失败）
# 重新用 raw_text 解析，更新非标准 YYYY-MM-DD 的 booking_date
try:
    import re as _re
    _db = get_db()
    _cur = _db.cursor()
    _cur.execute("SELECT id, raw_text, booking_date, booking_time, meal_time FROM orders")
    _fixed = 0
    _fixed_mt = 0
    for _r in _cur.fetchall():
        _bd = _r["booking_date"] or ""
        _needs_update = False
        _new_bd = _bd
        _new_bt = _r["booking_time"] or ""
        _new_mt = _r["meal_time"] or ""
        # 标准 YYYY-MM-DD 且不等于"待定"之类垃圾值则跳过日期解析
        if not _re.match(r"^\d{4}-\d{2}-\d{2}$", _bd) and _r["raw_text"]:
            from parser import parse_order_text as _pot
            _parsed = _pot(_r["raw_text"])
            _new_bd = _parsed.get("booking_date", "")
            _new_bt = _parsed.get("booking_time", "") or _new_bt
            _new_mt = _parsed.get("meal_time", "") or _new_mt
            if _new_bd:
                _needs_update = True
        # 同时清洗历史脏 meal_time（即使日期正常）
        if _r["raw_text"]:
            from parser import parse_order_text as _pot, _clean_meal_time as _cmt
            _parsed2 = _pot(_r["raw_text"])
            _cleaned_mt = _parsed2.get("meal_time", "")
            if _cleaned_mt and _cleaned_mt != _r["meal_time"]:
                _new_mt = _cleaned_mt
                _needs_update = True
        if _needs_update:
            _cur.execute("UPDATE orders SET booking_date=?, booking_time=?, meal_time=? WHERE id=?",
                         (_new_bd, _new_bt, _new_mt, _r["id"]))
            _fixed += 1
            if _new_mt != (_r["meal_time"] or ""):
                _fixed_mt += 1
    if _fixed:
        _db.commit()
        print(f"[migrate] 已修复 {_fixed} 个订单（其中 {_fixed_mt} 个 meal_time 被清洗）")
    _db.close()
except Exception as _e:
    print(f"[migrate] 修复订单 booking_date/meal_time 跳过: {_e}")

# 迁移：7-8人餐 把"韩式蘸酱"替换为"酸辣烤肉汁"，与其他套餐小料顺序一致
try:
    _db = get_db()
    _cur = _db.cursor()
    _cur.execute("SELECT id FROM ingredients WHERE name='韩式蘸酱'")
    _hjj = _cur.fetchone()
    _cur.execute("SELECT id FROM ingredients WHERE name='酸辣烤肉汁'")
    _slj = _cur.fetchone()
    if _hjj and _slj:
        _hjj_id = _hjj["id"]
        _slj_id = _slj["id"]
        # 7-8人餐的韩式蘸酱关联改为酸辣烤肉汁（per_package 调整为50，与其他套餐一致）
        _cur.execute("""UPDATE package_ingredients
                        SET ingredient_id=?, per_package=50
                        WHERE ingredient_id=? AND package_id IN (SELECT id FROM packages WHERE name='7-8人餐')""",
                     (_slj_id, _hjj_id))
        _affected = _cur.rowcount
        if _affected:
            _db.commit()
            print(f"[migrate] 已把 7-8人餐 的韩式蘸酱({_affected}条)替换为酸辣烤肉汁")
    _db.close()
except Exception as _e:
    print(f"[migrate] 替换韩式蘸酱跳过: {_e}")


@app.before_request
def before():
    g.db = get_db()


@app.teardown_request
def teardown(exc):
    db = getattr(g, "db", None)
    if db is not None:
        db.close()


# 数据变更后自动备份到 GitHub（POST/PUT/DELETE 且非备份接口）
@app.after_request
def auto_backup(response):
    if request.method in ("POST", "PUT", "DELETE") and "/api/backup" not in request.path:
        try:
            from db_backup import backup_async
            backup_async()
        except Exception:
            pass
    return response


@app.route("/api/backup", methods=["POST"])
def manual_backup():
    """手动触发备份"""
    from db_backup import backup_to_github
    ok = backup_to_github()
    if ok:
        return jsonify({"ok": True, "msg": "备份成功"})
    return jsonify({"ok": False, "msg": "备份失败（未配置 GITHUB_TOKEN 或网络错误）"})


@app.route("/api/backup/status", methods=["GET"])
def backup_status():
    """查询备份配置状态"""
    from db_backup import GITHUB_TOKEN, _last_backup
    import time as _time
    return jsonify({
        "configured": bool(GITHUB_TOKEN),
        "last_backup": _last_backup,
        "last_backup_ago": f"{int(_time.time() - _last_backup)}s" if _last_backup else None,
    })


# ===== 页面 =====
@app.route("/")
def index():
    return render_template("dashboard.html")  # 默认进数据面板


@app.route("/dashboard")
def dashboard_page():
    return render_template("dashboard.html")


@app.route("/config")
def config_page():
    return render_template("config.html")


@app.route("/orders")
def orders_page():
    return render_template("orders.html")


@app.route("/orders/<int:oid>")
def order_detail_page(oid):
    return render_template("order_detail.html", oid=oid)


# ===== 出餐明细单（客户小票：可打印/存PDF/导出图片）=====
@app.route("/orders/<int:oid>/ticket")
def order_ticket(oid):
    """从订单生成出餐明细单（58mm 小票样式，给客户核对菜品用）"""
    import html as _html
    import json as _json
    db = g.db
    cur = db.cursor()
    cur.execute("SELECT * FROM orders WHERE id = ?", (oid,))
    r = cur.fetchone()
    if not r:
        return "订单不存在", 404
    o = dict(r)
    cur.execute("""
        SELECT op.quantity, op.people, p.name as pkg_name, p.price
        FROM order_packages op LEFT JOIN packages p ON op.package_id = p.id
        WHERE op.order_id = ?
        ORDER BY op.id
    """, (oid,))
    pkgs = [dict(x) for x in cur.fetchall()]

    def esc(s):
        return _html.escape(str(s if s is not None else ''))

    shop = "刘和牛户外烤肉"
    no = "#%d" % oid
    time_str = (o.get("pickup_time") or o.get("meal_time") or
                ((o.get("booking_date") or "") + " " + (o.get("booking_time") or "")).strip() or
                "时间待定")
    contact = (o.get("contact_name") or "").strip()
    phone = (o.get("contact_phone") or "").strip()
    customer = (contact + (" " + phone if phone else "")).strip() or "—"
    address = (o.get("address") or "").strip()
    note = (o.get("note") or "").strip()

    calc_total = 0.0
    item_rows = ""
    data_items = []
    for p in pkgs:
        qty = int(p.get("quantity") or 1)
        price = float(p.get("price") or 0)
        line = round(price * qty, 2)
        calc_total += line
        price_str = ("¥%.2f" % line) if price > 0 else "—"
        item_rows += ('<div class="item"><div class="ln"><span class="nm">%s</span>'
                      '<span>x%d</span><span>%s</span></div></div>'
                      % (esc(p.get("pkg_name")), qty, price_str))
        data_items.append({"n": p.get("pkg_name") or "", "q": qty, "ps": price_str})
    if not item_rows:
        item_rows = '<div class="item">（本单未选套餐）</div>'
    if o.get("amount"):
        total_str = "¥%.2f" % float(o["amount"])
    elif calc_total > 0:
        total_str = "¥%.2f" % round(calc_total, 2)
    else:
        total_str = "—"
    deposit = float(o.get("deposit") or 0)

    # 优惠/加收：总额与套餐标价之和对不上时，用这一行把账做平
    discount = float(o.get("discount") or 0)
    adj_line = ""
    if abs(discount) > 0.001 and calc_total > 0:
        adj_line = ('<div class="item subst"><div class="ln"><span class="nm">套餐小计</span>'
                    '<span></span><span>¥%.2f</span></div></div>' % round(calc_total, 2))
        if discount > 0:
            adj_line += ('<div class="item subst"><div class="ln"><span class="nm">优惠</span>'
                         '<span></span><span>-¥%.2f</span></div></div>' % discount)
        else:
            adj_line += ('<div class="item subst"><div class="ln"><span class="nm">加收</span>'
                         '<span></span><span>+¥%.2f</span></div></div>' % abs(discount))

    data = {
        "shop": shop, "no": no, "time": time_str,
        "customer": customer, "address": address,
        "items": data_items, "total": total_str,
        "deposit": ("¥%.2f" % deposit) if deposit > 0 else "",
        "note": note,
        "subtotal": ("¥%.2f" % round(calc_total, 2)) if calc_total > 0 else "",
        "discount_num": discount,
        "discount_str": (("-¥%.2f" % discount) if discount > 0 else ("+¥%.2f" % abs(discount))) if abs(discount) > 0.001 else "",
    }

    html = """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>出餐明细单 __NO__</title>
<style>
* { box-sizing: border-box; }
body { margin:0; background:#eef1f5; color:#1f2329; font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif; font-size:14px; }
.toolbar { text-align:center; padding:14px 10px 4px; }
.toolbar button { margin:4px 6px; padding:10px 22px; border:none; border-radius:8px; font-size:15px; font-weight:700; cursor:pointer; }
.b-print { background:#c0392b; color:#fff; }
.b-png { background:#27ae60; color:#fff; }
.hint { text-align:center; font-size:12px; color:#8a97a5; padding:2px 10px 8px; }
#pngBox { text-align:center; padding:6px 0 20px; }
#pngBox img { width:100%; max-width:260px; border:1px solid #d4dae3; border-radius:6px; background:#fff; }
#pngHint { display:none; text-align:center; color:#27ae60; font-size:13px; font-weight:600; padding:4px 0 12px; }
.wrap { display:flex; justify-content:center; padding:10px 8px 30px; }
.receipt { width:58mm; background:#fff; color:#000; padding:2mm 3mm; font-family:"Courier New","PingFang SC",monospace; font-size:12px; line-height:1.45; }
.receipt .c { text-align:center; }
.receipt .shop { font-size:15px; font-weight:bold; letter-spacing:1px; }
.receipt .title { font-size:13px; border-top:1px dashed #000; border-bottom:1px dashed #000; padding:3px 0; margin:4px 0; text-align:center; letter-spacing:2px; }
.receipt .meta { display:flex; justify-content:space-between; font-size:11px; }
.receipt .sep { border-top:1px dashed #000; margin:4px 0; }
.receipt .item { margin:2px 0; }
.receipt .item .nm { word-break:break-all; }
.receipt .item .ln { display:flex; justify-content:space-between; gap:6px; }
.receipt .item.subst { color:#666; font-size:13px; }
.receipt .grand { font-size:14px; font-weight:bold; display:flex; justify-content:space-between; }
.receipt .ft { text-align:center; font-size:10.5px; margin-top:6px; color:#333; }
.receipt .barcode { text-align:center; font-family:"Courier New",monospace; font-size:10px; letter-spacing:1px; margin-top:2px; }
@media print {
  body { background:#fff; }
  .toolbar, .hint, #pngBox, #pngHint { display:none !important; }
  .wrap { padding:0; }
  @page { size: 58mm auto; margin: 2mm; }
}
</style>
</head>
<body>
<div class="toolbar">
  <button class="b-print" onclick="window.print()">🖨️ 打印 / 存为PDF</button>
  <button class="b-png" onclick="exportPNG()">📷 导出图片</button>
  <button class="b-edit" onclick="location.href='/orders/__OID__'" style="background:#f0f2f5;color:#1f2329">✏️ 编辑本单</button>
</div>
<div class="hint">手机端：点「导出图片」后在下方长按图片即可保存/转发</div>
<div id="pngHint">✅ 图片已生成，长按下方图片可保存或转发</div>
<div id="pngBox"></div>
<div class="wrap">
<div class="receipt" id="rcpt">
<div class="c shop">__SHOP__</div>
<div class="c" style="font-size:11px">出 餐 明 细 单</div>
<div class="title">*** 出餐明细单 ***</div>
<div class="meta"><span>单号：__NO__</span><span>__TIME__</span></div>
__META_EXTRA__
__ITEMS__
<div class="sep"></div>
__ADJ_LINE__
<div class="grand"><span>合计</span><span>__TOTAL__</span></div>
__DEPOSIT_LINE__
__NOTE_LINE__
<div class="ft">谢谢惠顾 · 请按小票核对菜品<br>祝您用餐愉快！</div>
<div class="barcode">__NO__</div>
</div>
</div>
<script>
var DATA = __DATA__;
function drawTicket(){
  var S = 2, W = 219;
  var lh = { c:18, title:22, sep:10, meta:16, grp:18, item:16, grand:20, note:16, ft:26, bar:16 };
  var L = [];
  L.push({t:'c', bold:true, size:15, text:DATA.shop});
  L.push({t:'c', size:11, text:'出 餐 明 细 单'});
  L.push({t:'title', text:'*** 出餐明细单 ***'});
  L.push({t:'meta', l:'单号：'+DATA.no, r:DATA.time});
  if (DATA.customer && DATA.customer !== '—') L.push({t:'meta', l:'客户：'+DATA.customer, r:''});
  if (DATA.address) L.push({t:'meta', l:'地址：'+DATA.address, r:''});
  L.push({t:'sep'});
  DATA.items.forEach(function(it){ L.push({t:'item', n:it.n, q:'x'+it.q, p:it.ps}); });
  L.push({t:'sep'});
  if (DATA.discount_num && DATA.subtotal) {
    L.push({t:'item', n:'套餐小计', q:'', p:DATA.subtotal});
    L.push({t:'item', n:(DATA.discount_num > 0 ? '优惠' : '加收'), q:'', p:DATA.discount_str});
  }
  L.push({t:'grand', l:'合计', r:DATA.total});
  if (DATA.deposit) L.push({t:'meta', l:'押金 '+DATA.deposit+'（离场退还）', r:''});
  if (DATA.note) { L.push({t:'sep'}); L.push({t:'note', text:'备注：'+DATA.note}); }
  L.push({t:'ft', text:'谢谢惠顾 · 请按小票核对菜品\\n祝您用餐愉快！'});
  L.push({t:'bar', text:DATA.no});
  var H = 10; L.forEach(function(l){ H += (lh[l.t] || 16); });
  var cv = document.createElement('canvas'); cv.width = W*S; cv.height = H*S;
  var ctx = cv.getContext('2d'); ctx.scale(S,S);
  ctx.fillStyle = '#fff'; ctx.fillRect(0,0,W,H); ctx.fillStyle = '#000'; ctx.textBaseline = 'top';
  var y = 8;
  function dash(x1,y1,x2,y2){ ctx.save(); ctx.setLineDash([3,2]); ctx.beginPath(); ctx.moveTo(x1,y1); ctx.lineTo(x2,y2); ctx.stroke(); ctx.restore(); }
  L.forEach(function(l){
    if (l.t==='c'){ ctx.textAlign='center'; ctx.font=(l.bold?'bold ':'')+(l.size||12)+'px "PingFang SC","Microsoft YaHei",monospace'; ctx.fillText(l.text, W/2, y); y += lh.c; }
    else if (l.t==='title'){ ctx.textAlign='center'; ctx.font='bold 13px "PingFang SC","Microsoft YaHei",monospace'; dash(4,y,W-4,y); y+=5; ctx.fillText(l.text, W/2, y); y += lh.title; }
    else if (l.t==='sep'){ dash(4,y,W-4,y); y += lh.sep; }
    else if (l.t==='meta'){ ctx.textAlign='left'; ctx.font='11px "PingFang SC","Microsoft YaHei",monospace'; ctx.fillText(l.l,4,y); ctx.textAlign='right'; ctx.fillText(l.r||'',W-4,y); ctx.textAlign='left'; y += lh.meta; }
    else if (l.t==='item'){ ctx.textAlign='left'; ctx.font='12px "PingFang SC","Microsoft YaHei",monospace'; ctx.fillText(l.n,4,y); ctx.textAlign='right'; ctx.fillText(l.q+'   '+l.p, W-4, y); ctx.textAlign='left'; y += lh.item; }
    else if (l.t==='grand'){ ctx.textAlign='left'; ctx.font='bold 14px "PingFang SC","Microsoft YaHei",monospace'; ctx.fillText(l.l,4,y); ctx.textAlign='right'; ctx.fillText(l.r,W-4,y); ctx.textAlign='left'; y += lh.grand; }
    else if (l.t==='note'){ ctx.textAlign='left'; ctx.font='12px "PingFang SC","Microsoft YaHei",monospace'; ctx.fillText(l.text,4,y); y += lh.note; }
    else if (l.t==='ft'){ ctx.textAlign='center'; ctx.font='10.5px "PingFang SC","Microsoft YaHei",monospace'; l.text.split('\\n').forEach(function(s){ ctx.fillText(s, W/2, y); y += 13; }); }
    else if (l.t==='bar'){ ctx.textAlign='center'; ctx.font='10px "Courier New",monospace'; ctx.fillText(l.text, W/2, y); y += lh.bar; }
  });
  return cv;
}
function exportPNG(){
  var url = drawTicket().toDataURL('image/png');
  document.getElementById('pngBox').innerHTML = '<img src="'+url+'" alt="出餐单图片">';
  document.getElementById('pngHint').style.display = 'block';
  try { var a = document.createElement('a'); a.href = url; a.download = '出餐单_'+DATA.no.replace('#','')+'.png'; document.body.appendChild(a); a.click(); a.remove(); } catch(e) {}
}
</script>
</body>
</html>"""

    meta_extra = ""
    if customer != "—":
        meta_extra += '<div class="meta"><span>客户：%s</span></div>' % esc(customer)
    if address:
        meta_extra += '<div class="meta"><span>地址：%s</span></div>' % esc(address)
    deposit_line = ""
    if deposit > 0:
        deposit_line = '<div class="meta" style="margin-top:2px"><span>押金 ¥%.2f（离场退还）</span></div>' % deposit
    note_line = ""
    if note:
        note_line = '<div class="sep"></div><div class="item">备注：%s</div>' % esc(note)

    html = (html
            .replace("__SHOP__", esc(shop))
            .replace("__NO__", esc(no))
            .replace("__OID__", str(oid))
            .replace("__TIME__", esc(time_str))
            .replace("__META_EXTRA__", meta_extra)
            .replace("__ITEMS__", item_rows)
            .replace("__ADJ_LINE__", adj_line)
            .replace("__TOTAL__", esc(total_str))
            .replace("__DEPOSIT_LINE__", deposit_line)
            .replace("__NOTE_LINE__", note_line)
            .replace("__DATA__", _json.dumps(data, ensure_ascii=False)))
    return html


@app.route("/api/orders/<int:oid>/detail")
def order_detail(oid):
    """订单详情：基本信息 + 备餐需求 + 工具借还 + 成本明细"""
    db = g.db
    cur = db.cursor()
    cur.execute("SELECT * FROM orders WHERE id = ?", (oid,))
    r = cur.fetchone()
    if not r:
        return jsonify({"ok": False, "msg": "订单不存在"}), 404
    o = dict(r)
    # 套餐
    cur.execute("""
        SELECT op.*, p.name as pkg_name, p.min_people, p.max_people
        FROM order_packages op LEFT JOIN packages p ON op.package_id = p.id
        WHERE op.order_id = ?
    """, (oid,))
    o["packages"] = [dict(x) for x in cur.fetchall()]
    # 备餐需求
    from calculator import calc_order_requirements, calc_prep_urgency
    req = calc_order_requirements(oid)
    o["requirements"] = req
    o["urgency"] = calc_prep_urgency(o)
    # 工具借还
    cur.execute("""
        SELECT tl.*, t.name as tool_name, t.cost
        FROM tool_loans tl LEFT JOIN tools t ON tl.tool_id = t.id
        WHERE tl.order_id = ?
    """, (oid,))
    o["tool_loans"] = [dict(x) for x in cur.fetchall()]
    return jsonify({"ok": True, "data": o})


@app.route("/history")
@app.route("/finance")
def finance_page():
    return render_template("finance.html")


@app.route("/prep")
def prep_page():
    return render_template("prep.html")


@app.route("/menu")
def menu_page():
    return render_template("menu.html")


# ===== 订单解析与入库 =====
@app.route("/api/parse", methods=["POST"])
def api_parse():
    """解析订单文本预览（不入库）。支持批量：用 --- 分隔多个订单"""
    data = request.get_json(force=True)
    raw = data.get("raw_text", "")
    if not raw.strip():
        return jsonify({"ok": False, "msg": "文本为空"}), 400
    import re as _re
    chunks = _re.split(r'\n[\-=]{3,}\n', raw.strip())
    chunks = [c.strip() for c in chunks if c.strip()]
    if len(chunks) > 1:
        previews = []
        for i, chunk in enumerate(chunks):
            parsed = preview_parse(chunk)
            parsed["_batch_idx"] = i + 1
            previews.append(parsed)
        return jsonify({"ok": True, "batch": True, "count": len(previews), "items": previews})
    parsed = preview_parse(raw)
    return jsonify({"ok": True, "data": parsed})


def _dup_key_of(parsed):
    """取用于「完全重复」判定的关键字段；关键信息缺失时返回 None（不做重复判断）"""
    bd = (parsed.get("booking_date") or "").strip()
    addr = (parsed.get("address") or "").strip()
    if not bd or not addr:
        return None
    return (bd, (parsed.get("booking_time") or "").strip(), addr,
            (parsed.get("contact_phone") or "").strip())


def _find_duplicate_order(cur, parsed):
    """找一条「完全重复」的现存订单（同日期+时间+地址+电话，且未取消、未删除）"""
    key = _dup_key_of(parsed)
    if not key:
        return None
    return cur.execute("""
        SELECT id, booking_date, booking_time, address, contact_name, contact_phone
        FROM orders
        WHERE deleted_at IS NULL AND status != 'cancelled'
          AND booking_date = ?
          AND COALESCE(booking_time,'') = ?
          AND COALESCE(address,'') = ?
          AND COALESCE(contact_phone,'') = ?
        ORDER BY id DESC LIMIT 1
    """, key).fetchone()


@app.route("/api/orders", methods=["POST"])
def create_order():
    """创建订单：接收 raw_text（单个或批量用 --- 分隔），自动解析入库。
    若发现一模一样的订单，返回 409 + duplicate 标记，由前端弹框让用户确认；
    用户确认后带 force=true 再提交一次即强制入库。"""
    data = request.get_json(force=True)
    raw = data.get("raw_text", "")
    if not raw.strip():
        return jsonify({"ok": False, "msg": "文本为空"}), 400

    # 批量：用 --- 或 === 分隔多个订单
    import re as _re
    chunks = _re.split(r'\n[\-=]{3,}\n', raw.strip())
    chunks = [c.strip() for c in chunks if c.strip()]

    db = g.db
    cur = db.cursor()
    parsed_chunks = [parse_order_text(c) for c in chunks]

    # ---- 重复订单检查（调用方已 force 则跳过）----
    if not data.get("force"):
        for pc in parsed_chunks:
            dup = _find_duplicate_order(cur, pc)
            if dup:
                d = dict(dup)
                where = " ".join(x for x in [d.get("booking_date") or "",
                                             d.get("booking_time") or "",
                                             d.get("address") or "",
                                             d.get("contact_name") or ""] if x)
                return jsonify({
                    "ok": False,
                    "duplicate": True,
                    "existing_id": d["id"],
                    "existing": d,
                    "msg": "已有一模一样的订单 #%s（%s）" % (d["id"], where),
                }), 409

    order_ids = []
    results = []
    for chunk, parsed in zip(chunks, parsed_chunks):
        matched = match_packages_in_db(parsed["packages"])
        cur.execute("""
            INSERT INTO orders
            (raw_text, booking_date, booking_time, address, contact_name,
             contact_phone, amount, deposit, meal_time, pickup_time, note, status)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,'pending')
        """, (
            chunk, parsed["booking_date"], parsed["booking_time"],
            parsed["address"], parsed["contact_name"], parsed["contact_phone"],
            parsed["amount"], parsed["deposit"], parsed["meal_time"],
            parsed["pickup_time"], parsed["note"],
        ))
        order_id = cur.lastrowid
        order_ids.append(order_id)
        for pk in matched:
            if not pk["package_id"]:
                continue
            cur.execute("""
                INSERT INTO order_packages
                (order_id, package_id, people, quantity)
                VALUES (?,?,?,?)
            """, (order_id, pk["package_id"], pk["people"], pk["quantity"]))
        results.append({"order_id": order_id, "name": parsed.get("contact_name", ""),
                        "date": parsed.get("booking_date", ""), "amount": parsed.get("amount", 0)})

    db.commit()
    if len(results) == 1:
        return jsonify({"ok": True, "order_id": order_ids[0]})
    return jsonify({"ok": True, "batch": True, "count": len(results), "orders": results})


@app.route("/api/orders")
def list_orders():
    db = g.db
    cur = db.cursor()
    purge_expired_deleted(db)          # 超过 7 天的回收站订单自动清掉
    show_deleted = request.args.get("deleted") == "1"
    if show_deleted:
        cur.execute("""
            SELECT id, booking_date, booking_time, address, contact_name,
                   contact_phone, amount, deposit, note, status, created_at, deleted_at
            FROM orders WHERE deleted_at IS NOT NULL ORDER BY id DESC LIMIT 100
        """)
    else:
        cur.execute("""
            SELECT id, booking_date, booking_time, address, contact_name,
                   contact_phone, amount, deposit, note, status, created_at
            FROM orders WHERE deleted_at IS NULL ORDER BY id DESC LIMIT 100
        """)
    orders = [dict(r) for r in cur.fetchall()]
    # 加套餐明细
    for o in orders:
        cur.execute("""
            SELECT op.people, op.quantity, p.name
            FROM order_packages op
            LEFT JOIN packages p ON op.package_id = p.id
            WHERE op.order_id = ?
        """, (o["id"],))
        o["packages"] = [dict(r) for r in cur.fetchall()]
    deleted_count = cur.execute(
        "SELECT COUNT(*) FROM orders WHERE deleted_at IS NOT NULL").fetchone()[0]
    return jsonify({"ok": True, "data": orders, "deleted_count": deleted_count})


PURGE_DAYS = 7


def purge_expired_deleted(db=None):
    """回收站里超过 7 天的订单彻底删除（含级联的套餐/借还/勾选）"""
    db = db or g.db
    try:
        db.execute("""
            DELETE FROM orders
            WHERE deleted_at IS NOT NULL
              AND deleted_at < datetime('now','localtime','-%d days')
        """ % PURGE_DAYS)
        db.commit()
    except Exception:
        pass


@app.route("/api/orders/<int:oid>/restore", methods=["POST"])
def restore_order(oid):
    """从回收站恢复订单"""
    db = g.db
    cur = db.execute("SELECT id FROM orders WHERE id=?", (oid,))
    if not cur.fetchone():
        return jsonify({"ok": False, "msg": "订单不存在"}), 404
    db.execute("UPDATE orders SET deleted_at=NULL WHERE id=?", (oid,))
    db.commit()
    return jsonify({"ok": True, "msg": "已恢复"})


@app.route("/api/orders/<int:oid>/purge", methods=["DELETE"])
def purge_order(oid):
    """彻底删除（不可恢复）"""
    db = g.db
    cur = db.execute("SELECT id FROM orders WHERE id=?", (oid,))
    if not cur.fetchone():
        return jsonify({"ok": False, "msg": "订单不存在"}), 404
    db.execute("DELETE FROM orders WHERE id=?", (oid,))
    db.commit()
    return jsonify({"ok": True, "msg": "已彻底删除"})


@app.route("/api/orders/<int:oid>/edit", methods=["PUT", "POST"])
def edit_order(oid):
    """编辑订单：可改 预约日期/时间、用餐时间、地址、联系人/电话、金额、备注。
    传了 packages 就整单替换套餐明细：[{package_id, people, quantity}, ...]"""
    data = request.get_json(force=True)
    db = g.db
    cur = db.cursor()
    if not cur.execute("SELECT id FROM orders WHERE id=?", (oid,)).fetchone():
        return jsonify({"ok": False, "msg": "订单不存在"}), 404

    def _txt(k):
        return (data.get(k) or "").strip()

    sets, params = [], []
    for k in ("booking_date", "booking_time", "meal_time", "note",
              "address", "contact_name", "contact_phone"):
        if k in data:
            sets.append("%s=?" % k)
            params.append(_txt(k))
    if "amount" in data:
        try:
            amt = float(data.get("amount") or 0)
        except (TypeError, ValueError):
            return jsonify({"ok": False, "msg": "金额格式不对"}), 400
        sets.append("amount=?")
        params.append(amt)
    if "discount" in data:
        try:
            disc = float(data.get("discount") or 0)
        except (TypeError, ValueError):
            return jsonify({"ok": False, "msg": "优惠格式不对"}), 400
        sets.append("discount=?")
        params.append(disc)
    if sets:
        params.append(oid)
        cur.execute("UPDATE orders SET %s WHERE id=?" % ", ".join(sets), params)

    if "packages" in data:
        rows = []
        for p in (data.get("packages") or []):
            try:
                pid = int(p.get("package_id"))
            except (TypeError, ValueError):
                continue
            try:
                qty = int(p.get("quantity") or 1)
            except (TypeError, ValueError):
                qty = 1
            if qty <= 0:
                continue
            pk = cur.execute("SELECT max_people FROM packages WHERE id=?", (pid,)).fetchone()
            if not pk:
                continue
            people = p.get("people")
            if people in (None, ""):
                people = pk["max_people"]
            rows.append((pid, int(people or 0), qty))
        cur.execute("DELETE FROM order_packages WHERE order_id=?", (oid,))
        for pid, people, qty in rows:
            cur.execute("INSERT INTO order_packages (order_id, package_id, people, quantity) "
                        "VALUES (?,?,?,?)", (oid, pid, people, qty))

    db.commit()
    return jsonify({"ok": True, "msg": "已保存"})


@app.route("/api/orders/<int:oid>", methods=["PATCH", "DELETE"])
def update_order_status(oid):
    db = g.db
    # 删除订单 → 进回收站（软删除），7 天内可恢复
    if request.method == "DELETE":
        cur = db.execute("SELECT id FROM orders WHERE id=?", (oid,))
        if not cur.fetchone():
            return jsonify({"ok": False, "msg": "订单不存在"}), 404
        db.execute("UPDATE orders SET deleted_at=datetime('now','localtime') WHERE id=?", (oid,))
        db.commit()
        return jsonify({"ok": True, "msg": "已放进回收站，7 天内可恢复"})
    data = request.get_json(force=True)
    status = data.get("status")
    if status not in ("pending", "preparing", "done", "cancelled"):
        return jsonify({"ok": False, "msg": "状态非法"}), 400
    db.execute("UPDATE orders SET status=? WHERE id=?", (status, oid))
    # 转备餐中 → 自动借出该订单所需工具
    if status == "preparing":
        borrow_tools(db, oid)
    db.commit()
    return jsonify({"ok": True})


def borrow_tools(db, order_id):
    """根据订单需求创建工具借出记录（若已存在则跳过）"""
    req = calc_order_requirements(order_id)
    cur = db.cursor()
    cur.execute("SELECT tool_id FROM tool_loans WHERE order_id=?", (order_id,))
    existing = {r["tool_id"] for r in cur.fetchall()}
    for t in req["tools"]:
        if t["id"] in existing:
            continue
        cur.execute("""
            INSERT INTO tool_loans (order_id, tool_id, quantity, status)
            VALUES (?,?,?,'borrowed')
        """, (order_id, t["id"], t["need"]))


@app.route("/api/orders/<int:oid>/requirements")
def order_requirements(oid):
    req = calc_order_requirements(oid)
    return jsonify({"ok": True, "data": req})


# ===== 库存面板 =====
@app.route("/api/dashboard")
def dashboard_api():
    return jsonify({"ok": True, "data": calc_dashboard()})


# ===== 库存调整 =====
@app.route("/api/stock/adjust", methods=["POST"])
def adjust_stock():
    """调整库存：item_type(ingredient/tool), item_id, delta, reason"""
    data = request.get_json(force=True)
    itype = data.get("item_type")
    iid = data.get("item_id")
    delta = float(data.get("delta", 0))
    reason = data.get("reason", "")
    if itype not in ("ingredient", "tool") or not iid:
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    db = g.db
    table = "ingredients" if itype == "ingredient" else "tools"
    db.execute(f"UPDATE {table} SET stock = stock + ? WHERE id = ?", (delta, iid))
    db.execute("""
        INSERT INTO stock_logs (item_type, item_id, delta, reason)
        VALUES (?,?,?,?)
    """, (itype, iid, delta, reason))
    db.commit()
    return jsonify({"ok": True})


# ===== 进货单解析 =====
@app.route("/api/stock/parse", methods=["POST"])
def api_stock_parse():
    """解析进货单文本，返回匹配预览"""
    from parser import parse_stock_invoice_text, match_inbound_to_db
    data = request.get_json(force=True)
    raw = data.get("raw_text", "")
    if not raw.strip():
        return jsonify({"ok": False, "msg": "文本为空"}), 400
    items = parse_stock_invoice_text(raw)
    matched = match_inbound_to_db(items)
    ok = sum(1 for x in matched if x["ok"])
    return jsonify({"ok": True, "items": matched, "count": len(matched), "matched": ok, "unmatched": len(matched) - ok})


@app.route("/api/stock/inbound", methods=["POST"])
def api_stock_inbound():
    """确认入库：接收 [{item_type, item_id, qty, unit_price, supplier, note}] 列表
    unit_price 有值时按加权平均更新 cost，并记 purchases 采购记录。
    """
    data = request.get_json(force=True)
    items = data.get("items", [])
    reason = data.get("reason", "进货入库")
    if not items:
        return jsonify({"ok": False, "msg": "无商品"}), 400
    db = g.db
    done = 0
    for it in items:
        iid = it.get("item_id")
        itype = it.get("item_type", "ingredient")
        qty = float(it.get("qty", 0))
        unit_price = float(it.get("unit_price", 0) or 0)
        supplier = it.get("supplier", "")
        note = it.get("note", "") or reason
        if not iid or qty <= 0:
            continue
        table = "ingredients" if itype == "ingredient" else "tools"
        # 成本更新规则（用户定制）：
        #   涨价（新单价 > 旧成本）→ 直接用新单价（保守，配餐不亏本）
        #   跌价或持平（新单价 ≤ 旧成本）→ 加权平均（平滑波动）
        if unit_price > 0:
            row = db.execute(f"SELECT stock, cost FROM {table} WHERE id = ?", (iid,)).fetchone()
            if row:
                old_stock = row["stock"] or 0
                old_cost = row["cost"] or 0
                new_stock = old_stock + qty
                if new_stock > 0:
                    # 加权平均（用于库存价值核算）
                    weighted_cost = round((old_stock * old_cost + qty * unit_price) / new_stock, 4)
                    # 涨价用新价，跌价/持平用加权
                    if unit_price > old_cost and old_cost > 0:
                        new_cost = round(unit_price, 4)  # 涨价→新价
                    else:
                        new_cost = weighted_cost        # 跌价/持平→加权
                    db.execute(f"UPDATE {table} SET stock = stock + ?, cost = ? WHERE id = ?", (qty, new_cost, iid))
                else:
                    db.execute(f"UPDATE {table} SET stock = stock + ? WHERE id = ?", (qty, iid))
            else:
                db.execute(f"UPDATE {table} SET stock = stock + ? WHERE id = ?", (qty, iid))
        else:
            db.execute(f"UPDATE {table} SET stock = stock + ? WHERE id = ?", (qty, iid))
        # 库存日志
        db.execute("""
            INSERT INTO stock_logs (item_type, item_id, delta, reason)
            VALUES (?,?,?,?)
        """, (itype, iid, qty, reason))
        # 采购记录（有单价才记）
        if unit_price > 0:
            db.execute("""
                INSERT INTO purchases (item_type, item_id, quantity, unit_price, total_cost, supplier, note)
                VALUES (?,?,?,?,?,?,?)
            """, (itype, iid, qty, unit_price, round(qty * unit_price, 2), supplier, note))
        done += 1
    db.commit()
    return jsonify({"ok": True, "done": done})


@app.route("/api/purchases")
def list_purchases():
    """采购记录列表"""
    db = g.db
    rows = db.execute("""
        SELECT p.*, CASE WHEN p.item_type='ingredient' THEN i.name ELSE t.name END as item_name,
               CASE WHEN p.item_type='ingredient' THEN i.unit ELSE '个' END as item_unit
        FROM purchases p
        LEFT JOIN ingredients i ON p.item_id = i.id
        LEFT JOIN tools t ON p.item_id = t.id
        ORDER BY p.purchased_at DESC
        LIMIT 200
    """).fetchall()
    return jsonify({"ok": True, "data": [dict(r) for r in rows]})


# ================= 库存管理（总览 / 流水 / 出库 / 盘点） =================

def _stock_status(stock, thr):
    """库存分档：out=缺货(≤0)，low=低库存(≤阈值)，ok=正常"""
    stock = float(stock or 0)
    thr = float(thr or 0)
    if stock <= 0:
        return "out"
    if thr > 0 and stock <= thr:
        return "low"
    return "ok"


@app.route("/api/stock/overview")
def stock_overview():
    """库存总览：KPI + 缺货/低库存清单 + 待出库订单 + 近7天出入库"""
    db = g.db
    cur = db.cursor()
    ings = [dict(r) for r in cur.execute(
        "SELECT id,name,unit,stock,threshold,cost,category FROM ingredients").fetchall()]
    tools = [dict(r) for r in cur.execute(
        "SELECT id,name,stock,threshold,cost FROM tools").fetchall()]

    def summarize(items, unit_default=""):
        out, low, val = [], [], 0.0
        for x in items:
            x["status"] = _stock_status(x.get("stock"), x.get("threshold"))
            val += float(x.get("stock") or 0) * float(x.get("cost") or 0)
            if x["status"] == "out":
                out.append(x)
            elif x["status"] == "low":
                low.append(x)
        return {"count": len(items), "out": len(out), "low": len(low),
                "ok": len(items) - len(out) - len(low),
                "value": round(val, 2), "out_items": out[:50], "low_items": low[:50]}

    ing_s = summarize(ings)
    tool_s = summarize(tools, "个")

    pending = [dict(r) for r in cur.execute("""
        SELECT o.id, o.booking_date, o.booking_time, o.address, o.contact_name, o.status
        FROM orders o
        WHERE o.deleted_at IS NULL AND o.status != 'cancelled'
          AND (o.stock_deducted_at IS NULL OR o.stock_deducted_at = '')
        ORDER BY o.booking_date, o.booking_time
        LIMIT 60
    """).fetchall()]
    deducted = cur.execute("""
        SELECT COUNT(*) FROM orders
        WHERE deleted_at IS NULL AND status != 'cancelled'
          AND stock_deducted_at IS NOT NULL AND stock_deducted_at != ''
    """).fetchone()[0]

    r = cur.execute("""SELECT COALESCE(SUM(CASE WHEN delta>0 THEN delta ELSE 0 END),0) as inq,
                              COALESCE(SUM(CASE WHEN delta<0 THEN -delta ELSE 0 END),0) as outq,
                              COUNT(*) as c
                       FROM stock_logs
                       WHERE created_at >= datetime('now','localtime','-7 days')""").fetchone()
    recent = {"in": round(r["inq"], 2), "out": round(r["outq"], 2), "count": r["c"]}

    return jsonify({"ok": True, "data": {
        "ingredients": ing_s, "tools": tool_s,
        "total_value": round(ing_s["value"] + tool_s["value"], 2),
        "pending_consume": pending, "deducted_count": deducted,
        "recent7": recent,
    }})


@app.route("/api/stock/logs")
def stock_logs_api():
    """库存流水（最近 N 条，可按 item_type 过滤）"""
    try:
        limit = int(request.args.get("limit", 200))
    except (TypeError, ValueError):
        limit = 200
    limit = max(1, min(limit, 1000))
    itype = request.args.get("item_type") or ""
    sql = """
        SELECT l.*,
               CASE WHEN l.item_type='ingredient' THEN i.name ELSE t.name END as item_name,
               CASE WHEN l.item_type='ingredient' THEN i.unit ELSE '个' END as item_unit
        FROM stock_logs l
        LEFT JOIN ingredients i ON l.item_type='ingredient' AND l.item_id=i.id
        LEFT JOIN tools t ON l.item_type='tool' AND l.item_id=t.id
    """
    params = []
    if itype in ("ingredient", "tool"):
        sql += " WHERE l.item_type = ? "
        params.append(itype)
    sql += " ORDER BY l.id DESC LIMIT ?"
    params.append(limit)
    rows = [dict(r) for r in g.db.execute(sql, params).fetchall()]
    return jsonify({"ok": True, "data": rows})


@app.route("/api/stock/order-needs/<int:oid>")
def stock_order_needs(oid):
    """预览某单出库会扣掉什么（不落库）"""
    from calculator import calc_order_requirements
    db = g.db
    o = db.execute("SELECT id,booking_date,booking_time,address,contact_name,status,"
                   "stock_deducted_at FROM orders WHERE id=? AND deleted_at IS NULL", (oid,)).fetchone()
    if not o:
        return jsonify({"ok": False, "msg": "订单不存在"}), 404
    req = calc_order_requirements(oid)
    items = []
    for it in req.get("ingredients", []):
        if float(it.get("need") or 0) <= 0:
            continue
        items.append({"item_type": "ingredient", "item_id": it["id"], "name": it["name"],
                      "need": round(float(it["need"]), 2), "unit": it.get("unit") or "",
                      "stock": it.get("stock", 0)})
    for it in req.get("tools", []):
        if float(it.get("need") or 0) <= 0:
            continue
        items.append({"item_type": "tool", "item_id": it["id"], "name": it["name"],
                      "need": round(float(it["need"]), 2), "unit": "个",
                      "stock": it.get("stock", 0)})
    return jsonify({"ok": True, "data": {"order": dict(o), "items": items}})


@app.route("/api/stock/consume", methods=["POST"])
def stock_consume():
    """按订单出库：把该单所需食材/工具从库存扣掉并记流水（ref=order:<id>）。
    同一单重复调用默认拒绝（幂等保护）；要重扣先撤销。"""
    data = request.get_json(force=True)
    oid = data.get("order_id")
    if not oid:
        return jsonify({"ok": False, "msg": "缺少订单号"}), 400
    db = g.db
    cur = db.cursor()
    o = cur.execute("SELECT * FROM orders WHERE id=? AND deleted_at IS NULL", (oid,)).fetchone()
    if not o:
        return jsonify({"ok": False, "msg": "订单不存在"}), 404
    if (o["stock_deducted_at"] or "").strip():
        return jsonify({"ok": False, "already": True,
                        "msg": "订单 #%s 已经出过库了（%s），如需重扣请先撤销"
                               % (oid, o["stock_deducted_at"])}), 409

    from calculator import calc_order_requirements
    req = calc_order_requirements(int(oid))
    moved, short = 0, []
    for itype, key, table, unit_default in (("ingredient", "ingredients", "ingredients", ""),
                                            ("tool", "tools", "tools", "个")):
        for it in req.get(key, []):
            need = round(float(it.get("need") or 0), 4)
            if need <= 0:
                continue
            row = cur.execute("SELECT stock FROM %s WHERE id=?" % table, (it["id"],)).fetchone()
            if not row:
                continue
            avail = float(row["stock"] or 0)
            if avail < need:
                short.append({"name": it["name"], "need": need, "stock": avail,
                              "unit": it.get("unit") or unit_default})
            cur.execute("UPDATE %s SET stock = stock - ? WHERE id=?" % table, (need, it["id"]))
            cur.execute("INSERT INTO stock_logs (item_type,item_id,delta,reason,ref) "
                        "VALUES (?,?,?,?,?)",
                        (itype, it["id"], -need, "订单#%s 出餐出库" % oid, "order:%s" % oid))
            moved += 1
    cur.execute("UPDATE orders SET stock_deducted_at=datetime('now','localtime') WHERE id=?", (oid,))
    db.commit()
    return jsonify({"ok": True, "moved": moved, "short": short})


@app.route("/api/stock/consume/undo", methods=["POST"])
def stock_consume_undo():
    """撤销某单的出库：把该单产生的流水逐条反向，并删除这些流水。"""
    data = request.get_json(force=True)
    oid = data.get("order_id")
    if not oid:
        return jsonify({"ok": False, "msg": "缺少订单号"}), 400
    db = g.db
    cur = db.cursor()
    ref = "order:%s" % oid
    rows = cur.execute("SELECT * FROM stock_logs WHERE ref=?", (ref,)).fetchall()
    if not rows:
        return jsonify({"ok": False, "msg": "该单没有出库记录"}), 404
    for r in rows:
        table = "ingredients" if r["item_type"] == "ingredient" else "tools"
        cur.execute("UPDATE %s SET stock = stock - ? WHERE id=?" % table, (r["delta"], r["item_id"]))
    cur.execute("DELETE FROM stock_logs WHERE ref=?", (ref,))
    cur.execute("UPDATE orders SET stock_deducted_at=NULL WHERE id=?", (oid,))
    db.commit()
    return jsonify({"ok": True, "reverted": len(rows)})


@app.route("/api/stock/stocktake", methods=["POST"])
def stock_stocktake():
    """盘点：传每项实际数量，差额自动调整库存并记流水。
    请求体：{"items":[{"item_type":"ingredient","item_id":1,"actual":1234.5}, ...]}"""
    data = request.get_json(force=True)
    items = data.get("items") or []
    if not items:
        return jsonify({"ok": False, "msg": "没有盘点数据"}), 400
    db = g.db
    cur = db.cursor()
    changed = []
    for it in items:
        itype = it.get("item_type")
        iid = it.get("item_id")
        if itype not in ("ingredient", "tool") or not iid:
            continue
        try:
            actual = float(it.get("actual"))
        except (TypeError, ValueError):
            continue
        table = "ingredients" if itype == "ingredient" else "tools"
        row = cur.execute("SELECT name, stock FROM %s WHERE id=?" % table, (iid,)).fetchone()
        if not row:
            continue
        before = float(row["stock"] or 0)
        delta = round(actual - before, 4)
        if abs(delta) < 0.0001:
            continue
        cur.execute("UPDATE %s SET stock=? WHERE id=?" % table, (actual, iid))
        cur.execute("INSERT INTO stock_logs (item_type,item_id,delta,reason,ref) "
                    "VALUES (?,?,?,?,?)",
                    (itype, iid, delta, "盘点调整", "stocktake"))
        changed.append({"name": row["name"], "before": before, "after": actual, "delta": delta})
    db.commit()
    return jsonify({"ok": True, "changed": changed, "count": len(changed)})


@app.route("/api/ingredients/<int:iid>/cost", methods=["POST"])
def update_ingredient_cost(iid):
    """手动修改食材/工具成本单价"""
    data = request.get_json(force=True)
    cost = float(data.get("cost", 0))
    itype = data.get("item_type", "ingredient")
    table = "ingredients" if itype == "ingredient" else "tools"
    db = g.db
    db.execute(f"UPDATE {table} SET cost = ? WHERE id = ?", (cost, iid))
    db.commit()
    return jsonify({"ok": True})


# ===== 配置管理 =====
@app.route("/api/ingredients", methods=["GET", "POST"])
def manage_ingredients():
    db = g.db
    if request.method == "GET":
        rows = db.execute("SELECT * FROM ingredients ORDER BY id").fetchall()
        items = [dict(r) for r in rows]
        # meat 细分 + 固定排序（复用 calculator.sort_ingredients）
        from calculator import sort_ingredients
        items = sort_ingredients(items)
        return jsonify({"ok": True, "data": items})
    data = request.get_json(force=True)
    db.execute("""
        INSERT INTO ingredients (name, unit, stock, threshold, cost, category)
        VALUES (?,?,?,?,?,?)
    """, (data["name"], data.get("unit", ""), data.get("stock", 0),
          data.get("threshold", 0), data.get("cost", 0), data.get("category", "other")))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/ingredients/<int:iid>", methods=["PUT"])
def update_ingredient(iid):
    data = request.get_json(force=True)
    db = g.db
    # 支持部分更新：只更新传入的字段（避免未传字段被置空）
    cur = db.execute("SELECT * FROM ingredients WHERE id=?", (iid,))
    existing = cur.fetchone()
    if not existing:
        return jsonify({"ok": False, "msg": "not found"}), 404
    existing = dict(existing)
    fields = ["name", "unit", "stock", "threshold", "cost", "category"]
    updates = {}
    for f in fields:
        if f in data:
            updates[f] = data[f]
        else:
            updates[f] = existing[f]
    db.execute("""
        UPDATE ingredients SET name=?, unit=?, stock=?, threshold=?, cost=?, category=?
        WHERE id=?
    """, (updates["name"], updates["unit"], updates["stock"], updates["threshold"],
          updates["cost"], updates["category"], iid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/ingredients/<int:iid>", methods=["DELETE"])
def del_ingredient(iid):
    db = g.db
    db.execute("DELETE FROM ingredients WHERE id=?", (iid,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/tools", methods=["GET", "POST"])
def manage_tools():
    db = g.db
    if request.method == "GET":
        rows = db.execute("SELECT * FROM tools").fetchall()
        from calculator import sort_tools
        return jsonify({"ok": True, "data": sort_tools([dict(r) for r in rows])})
    data = request.get_json(force=True)
    db.execute("INSERT INTO tools (name, stock, threshold, cost) VALUES (?,?,?,?)",
              (data["name"], data.get("stock", 0), data.get("threshold", 0), data.get("cost", 0)))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/tools/<int:tid>", methods=["PUT", "DELETE"])
def update_tool(tid):
    db = g.db
    if request.method == "DELETE":
        db.execute("DELETE FROM tools WHERE id=?", (tid,))
    else:
        data = request.get_json(force=True)
        db.execute("UPDATE tools SET name=?, stock=?, threshold=? WHERE id=?",
                  (data["name"], data.get("stock", 0), data.get("threshold", 0), tid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/packages", methods=["GET", "POST"])
def manage_packages():
    db = g.db
    if request.method == "GET":
        cur = db.cursor()
        cur.execute("SELECT * FROM packages ORDER BY min_people")
        pkgs = [dict(r) for r in cur.fetchall()]
        for p in pkgs:
            cur.execute("""
                SELECT pi.id, pi.per_package, pi.portion_count,
                       i.name as ing_name, i.id as ing_id, i.unit, i.category
                FROM package_ingredients pi
                JOIN ingredients i ON pi.ingredient_id = i.id
                WHERE pi.package_id = ? AND pi.cost_only = 0
                ORDER BY i.category, i.name
            """, (p["id"],))
            p["ingredients"] = [dict(r) for r in cur.fetchall()]
            cur.execute("""
                SELECT pt.id, pt.per_package, t.name as tool_name, t.id as tool_id
                FROM package_tools pt
                JOIN tools t ON pt.tool_id = t.id
                WHERE pt.package_id = ?
            """, (p["id"],))
            from calculator import sort_tools as _st
            p["tools"] = _st([dict(r) for r in cur.fetchall()], "per_package", "tool_name")
        return jsonify({"ok": True, "data": pkgs})

    data = request.get_json(force=True)
    cur = db.cursor()
    # 新规则：接收 base_price（菜品价），price 由 base_price + delivery_fee 自动算
    base_price = float(data.get("base_price", data.get("price", 0)))
    df_row = db.execute("SELECT CAST(value AS REAL) as df FROM settings WHERE key='delivery_fee'").fetchone()
    df = df_row["df"] if df_row else 0
    total_price = base_price + df
    cur.execute("""
        INSERT INTO packages (name, min_people, max_people, base_price, price, service_type, is_team)
        VALUES (?,?,?,?,?,?,?)
    """, (data["name"], data["min_people"], data["max_people"],
          base_price, total_price, data.get("service_type", "外送"),
          data.get("is_team", 0)))
    pid = cur.lastrowid
    # 关联食材
    for ing in data.get("ingredients", []):
        cur.execute("""
            INSERT INTO package_ingredients (package_id, ingredient_id, per_package, portion_count)
            VALUES (?,?,?,?)
        """, (pid, ing["ingredient_id"], ing["per_package"], ing.get("portion_count", 1)))
    # 关联工具
    for t in data.get("tools", []):
        cur.execute("""
            INSERT INTO package_tools (package_id, tool_id, per_package)
            VALUES (?,?,?)
        """, (pid, t["tool_id"], t.get("per_package", 1)))
    db.commit()
    return jsonify({"ok": True, "id": pid})


@app.route("/api/packages/<int:pid>", methods=["DELETE"])
def del_package(pid):
    db = g.db
    db.execute("DELETE FROM packages WHERE id=?", (pid,))
    db.commit()
    return jsonify({"ok": True})


# ===== 餐标配餐器 =====
@app.route("/api/catering/suggest", methods=["POST"])
def catering_suggest():
    """餐标配餐器：输入人数+餐标+人工成本，返回建议套餐方案"""
    from calculator import catering_suggest as _cs
    data = request.get_json(force=True)
    people = int(data.get("people", 0))
    total = float(data.get("total_budget", 0))
    per_person = data.get("per_person")
    kitchen_labor = float(data.get("kitchen_labor", 0))
    delivery_labor = float(data.get("delivery_labor", 0))
    if per_person is not None:
        per_person = float(per_person)
    result = _cs(people, total, per_person, kitchen_labor, delivery_labor)
    return jsonify(result)


@app.route("/api/catering/save", methods=["POST"])
def catering_save():
    """把配餐方案保存为定制套餐"""
    data = request.get_json(force=True)
    name = data.get("name", "").strip()
    people = int(data.get("people", 0))
    total = float(data.get("total_budget", 0))
    kitchen_labor = float(data.get("kitchen_labor", 0))
    delivery_labor = float(data.get("delivery_labor", 0))
    ings = data.get("ingredients", [])
    tools = data.get("tools", [])
    if not name or people <= 0:
        return jsonify({"ok": False, "msg": "套餐名和人数必填"}), 400

    db = g.db
    cur = db.cursor()
    # base_price = 餐标 - 配送费
    df_row = db.execute("SELECT CAST(value AS REAL) as df FROM settings WHERE key='delivery_fee'").fetchone()
    df = df_row["df"] if df_row else 0
    base_price = max(0, total - df)
    total_price = base_price + df
    cur.execute("""
        INSERT INTO packages (name, min_people, max_people, base_price, price, service_type, is_team, kitchen_labor_cost, delivery_labor_cost)
        VALUES (?,?,?,?,?,?,?,?,?)
    """, (name, people, people, base_price, total_price, "搭建", 0, kitchen_labor, delivery_labor))
    pid = cur.lastrowid
    for ing in ings:
        cur.execute("""
            INSERT INTO package_ingredients (package_id, ingredient_id, per_package, portion_count, cost_only)
            VALUES (?,?,?,?,?)
        """, (pid, ing["ing_id"], ing["per_package"],
              ing.get("portion_count", 1), ing.get("cost_only", 0)))
    for t in tools:
        cur.execute("""
            INSERT INTO package_tools (package_id, tool_id, per_package)
            VALUES (?,?,?)
        """, (pid, t["id"], t.get("per_package", 1)))
    db.commit()
    return jsonify({"ok": True, "id": pid})


# 修改套餐里某食材的配量(每份克数 + 份数)
@app.route("/api/packages/<int:pid>/ingredient/<int:pi_id>", methods=["PUT", "DELETE"])
def update_pkg_ingredient(pid, pi_id):
    db = g.db
    cur = db.cursor()
    if request.method == "DELETE":
        cur.execute("DELETE FROM package_ingredients WHERE id=? AND package_id=?", (pi_id, pid))
        db.commit()
        return jsonify({"ok": True})
    data = request.get_json(force=True)
    size = float(data.get("portion_size", 0))
    count = float(data.get("portion_count", 1))
    total = size * count
    cur.execute("""
        UPDATE package_ingredients SET per_package=?, portion_count=?
        WHERE id=? AND package_id=?
    """, (total, count, pi_id, pid))
    db.commit()
    return jsonify({"ok": True, "per_package": total, "portion_count": count})


# 新增套餐食材
@app.route("/api/packages/<int:pid>/ingredient", methods=["POST"])
def add_pkg_ingredient(pid):
    db = g.db
    cur = db.cursor()
    data = request.get_json(force=True)
    ing_id = data.get("ingredient_id")
    size = float(data.get("portion_size", 0))
    count = float(data.get("portion_count", 1))
    total = size * count
    if not ing_id:
        return jsonify({"ok": False, "msg": "缺少食材ID"}), 400
    cur.execute("SELECT id FROM package_ingredients WHERE package_id=? AND ingredient_id=?", (pid, ing_id))
    if cur.fetchone():
        return jsonify({"ok": False, "msg": "该食材已存在，请直接修改"}), 400
    cur.execute("""
        INSERT INTO package_ingredients (package_id, ingredient_id, per_package, portion_count)
        VALUES (?,?,?,?)
    """, (pid, ing_id, total, count))
    db.commit()
    return jsonify({"ok": True, "id": cur.lastrowid})


# 修改套餐工具配量
@app.route("/api/packages/<int:pid>/tool/<int:pt_id>", methods=["PUT", "DELETE"])
def update_pkg_tool(pid, pt_id):
    db = g.db
    cur = db.cursor()
    if request.method == "DELETE":
        cur.execute("DELETE FROM package_tools WHERE id=? AND package_id=?", (pt_id, pid))
        db.commit()
        return jsonify({"ok": True})
    data = request.get_json(force=True)
    per = float(data.get("per_package", 1))
    cur.execute("UPDATE package_tools SET per_package=? WHERE id=? AND package_id=?", (per, pt_id, pid))
    db.commit()
    return jsonify({"ok": True})


# 导出菜单原数据(JSON / CSV / Excel)
@app.route("/api/menu/export")
def export_menu():
    import csv, io, json as _json
    fmt = request.args.get("format", "json")
    pid = request.args.get("pid")  # 单套餐导出
    cur = g.db.cursor()
    sql = "SELECT * FROM packages ORDER BY min_people"
    params = ()
    if pid:
        sql = "SELECT * FROM packages WHERE id=? ORDER BY min_people"
        params = (int(pid),)
    cur.execute(sql, params)
    pkgs = [dict(r) for r in cur.fetchall()]
    for p in pkgs:
        cur.execute("""
            SELECT i.name as ingredient, i.unit, i.category, pi.per_package, pi.portion_count
            FROM package_ingredients pi
            JOIN ingredients i ON pi.ingredient_id = i.id
            WHERE pi.package_id = ?
            ORDER BY i.category, i.name
        """, (p["id"],))
        p["ingredients"] = [dict(r) for r in cur.fetchall()]
        cur.execute("""
            SELECT t.name as tool, pt.per_package
            FROM package_tools pt
            JOIN tools t ON pt.tool_id = t.id
            WHERE pt.package_id = ?
        """, (p["id"],))
        from calculator import sort_tools as _st
        p["tools"] = _st([dict(r) for r in cur.fetchall()], "per_package", "tool")

    fname = f"menu_{pkgs[0]['name']}" if len(pkgs)==1 else "menu_all"

    if fmt == "csv":
        buf = io.StringIO()
        buf.write("\ufeff")  # BOM for Excel
        w = csv.writer(buf)
        w.writerow(["套餐", "人数范围", "类别", "食材/工具", "单位", "每份克数", "份数", "总量", "基价", "总价"])
        for p in pkgs:
            people = f"{p['min_people']}-{p['max_people']}人"
            for ing in p["ingredients"]:
                size = ing["per_package"] / max(1, ing["portion_count"])
                w.writerow([p["name"], people, ing["category"], ing["ingredient"],
                           ing["unit"], round(size, 2), ing["portion_count"],
                           ing["per_package"], p["base_price"], p["price"]])
            for t in p["tools"]:
                w.writerow([p["name"], people, "tool", t["tool"], "个", "", "",
                            t["per_package"], p["base_price"], p["price"]])
        resp = app.response_class(buf.getvalue(), mimetype="text/csv")
        resp.headers["Content-Disposition"] = f"attachment; filename={fname}.csv"
        return resp
    return jsonify({"ok": True, "data": pkgs})


# 导出 Excel(xlsx)
@app.route("/api/menu/export/xlsx")
def export_menu_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"ok": False, "msg": "需要 openpyxl: pip install openpyxl"}), 500
    from calculator import (
        CATEGORY_ORDER, CATEGORY_LABEL, _subcategorize_meat,
        PACKAGING_CATEGORY, UTENSIL_CATEGORY, sort_tools, _get_fixed_sort_index
    )
    pid = request.args.get("pid")
    cur = g.db.cursor()
    if pid:
        cur.execute("SELECT * FROM packages WHERE id=?", (int(pid),))
    else:
        cur.execute("SELECT * FROM packages ORDER BY min_people")
    pkgs = [dict(r) for r in cur.fetchall()]
    for p in pkgs:
        cur.execute("""
            SELECT i.name as ingredient, i.unit, i.category, pi.per_package, pi.portion_count
            FROM package_ingredients pi
            JOIN ingredients i ON pi.ingredient_id = i.id
            WHERE pi.package_id = ? AND pi.cost_only = 0
        """, (p["id"],))
        p["ingredients"] = [dict(r) for r in cur.fetchall()]
        cur.execute("""
            SELECT t.name as tool, pt.per_package
            FROM package_tools pt
            JOIN tools t ON pt.tool_id = t.id
            WHERE pt.package_id = ?
        """, (p["id"],))
        p["tools"] = sort_tools([dict(r) for r in cur.fetchall()], "per_package", "tool")

    # 与网页菜单完全一致的分类顺序与标签
    # 食材分类顺序：beef→pork→chicken→vegetable→side→sauce→drink→staple→packaging→utensil→other
    # 工具放最后
    EXPORT_CAT_ORDER = [
        "beef", "pork", "chicken", "vegetable", "side", "sauce", "drink",
        "staple", "packaging", "utensil", "other", "tool"
    ]
    cat_colors = {
        'beef': '8E1E1A', 'pork': 'C75D3E', 'chicken': 'C9962B',
        'vegetable': '3A8A3A', 'side': '5A8A3A', 'sauce': '7A5A3A',
        'drink': '666666', 'staple': '666666',
        'packaging': 'D35400', 'utensil': '8E44AD', 'other': '666666',
        'tool': '4A4A4A'
    }
    cat_labels = {
        'beef': '🥩 牛肉', 'pork': '🥓 猪肉', 'chicken': '🍗 鸡肉',
        'vegetable': '🥬 素菜', 'side': '🥗 小菜', 'sauce': '🧂 小料',
        'drink': '🎁 赠品', 'staple': '🍚 主食',
        'packaging': '📦 食材包装', 'utensil': '🍱 客户餐具',
        'other': '📦 其他', 'tool': '🔧 工具'
    }

    # 食材内部固定排序（与网页 FIXED_ORDER 一致）
    ing_fixed_order = [
        ['beef','肥牛'],['beef','拌牛肉'],['beef','牛肋条'],['beef','牛骰子'],['beef','小肠'],
        ['pork','五花肉'],['pork','小香猪'],['pork','松板肉'],['pork','风味肠'],['pork','梅花肉'],
        ['chicken','鸡尖'],['chicken','郡肝'],['chicken','鸡腿肉'],['chicken','鸡翅根'],
        ['chicken','掌中宝'],['chicken','鸡脚筋'],
        ['vegetable','生菜'],['vegetable','豆腐'],['vegetable','西葫芦'],
        ['vegetable','土豆'],['vegetable','馒头'],['vegetable','韭菜'],
        ['vegetable','杏鲍菇'],['vegetable','洋葱'],
        ['side','海带丝'],['side','辣椒段'],['side','辣白菜'],['side','蒜片'],
        ['sauce','川香'],['sauce','五香'],['sauce','酸辣'],
        ['drink','应季水果'],['drink','可乐'],['drink','雪碧'],
    ]
    def ing_sort_idx(name, cat):
        for i, (c, kw) in enumerate(ing_fixed_order):
            if c == cat and kw in (name or ''):
                return i
        return 999

    wb = Workbook()
    ws = wb.active
    ws.title = "菜单" if len(pkgs) > 1 else pkgs[0]["name"]
    # 与网页一致的表头（名称→份数→每份数量→单位→总量→保存→删除）
    headers = ["名称", "份数", "每份数量", "单位", "总量", "保存", "删除"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="3A2416")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(border_style="thin", color="FFFFFF"),
            right=Side(border_style="thin", color="FFFFFF"),
            top=Side(border_style="thin", color="FFFFFF"),
            bottom=Side(border_style="thin", color="FFFFFF"),
        )
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", indent=1)

    for p in pkgs:
        # 套餐标题行（合并 7 列）
        ws.append([f"{p['name']}  ·  {p['min_people']}-{p['max_people']}人  ·  基价¥{p['base_price']}  总价¥{p['price']}", "", "", "", "", "", ""])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor="8E1E1A")
            cell.font = Font(bold=True, color="FFFFFF", size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 食材按分类细分 + 固定排序
        by_cat = {}
        for ing in p["ingredients"]:
            raw_cat = ing["category"] or "other"
            # meat 细分为 beef/pork/chicken（与网页 subcategorize 一致）
            cat = _subcategorize_meat(ing["ingredient"], raw_cat)
            by_cat.setdefault(cat, []).append(ing)
        for cat in EXPORT_CAT_ORDER:
            items = by_cat.get(cat, [])
            if not items:
                continue
            # 分类内按固定顺序排序
            items.sort(key=lambda x: (ing_sort_idx(x["ingredient"], cat), x["ingredient"]))
            color = cat_colors.get(cat, "666666")
            label = cat_labels.get(cat, cat)
            # 分类标题行（与网页 cat-row 一致，本身是 7 列）
            ws.append([f"{label} · {len(items)}项", "份数", "每份数量", "单位", "总量", "保存", "删除"])
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(ws.max_row, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
            # 食材行（列序：名称→份数→每份数量→单位→总量→空→空）
            for ing in items:
                size = ing["per_package"] / max(1, ing["portion_count"])
                ws.append([
                    ing["ingredient"], ing["portion_count"], round(size, 2),
                    ing["unit"], ing["per_package"], "", ""
                ])
                for cell in ws[ws.max_row]:
                    cell.border = border
                    cell.alignment = center_align
                ws.cell(ws.max_row, 1).alignment = left_align
                ws.cell(ws.max_row, 5).font = Font(bold=True, color="8E1E1A")
        # 工具分类
        if p["tools"]:
            ws.append([f"🔧 工具 · {len(p['tools'])}项", "份数", "每份数量", "单位", "总量", "保存", "删除"])
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="4A4A4A")
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(ws.max_row, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for t in p["tools"]:
                ws.append([
                    t["tool"], 1, t["per_package"], "个", t["per_package"], "", ""
                ])
                for cell in ws[ws.max_row]:
                    cell.border = border
                    cell.alignment = center_align
                ws.cell(ws.max_row, 1).alignment = left_align
                ws.cell(ws.max_row, 5).font = Font(bold=True, color="4A4A4A")
        # 套餐间空行
        ws.append([])
    # 列宽（与网页列对齐）
    widths = [22, 8, 12, 8, 10, 8, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # 冻结表头
    ws.freeze_panes = "A2"
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    fname = f"menu_{pkgs[0]['name']}" if len(pkgs)==1 else "menu_all"
    resp = app.response_class(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["Content-Disposition"] = f"attachment; filename={fname}.xlsx"
    return resp


@app.route("/api/dishes", methods=["GET", "POST"])
def manage_dishes():
    db = g.db
    if request.method == "GET":
        rows = db.execute("SELECT * FROM dishes ORDER BY id").fetchall()
        return jsonify({"ok": True, "data": [dict(r) for r in rows]})
    data = request.get_json(force=True)
    db.execute("INSERT INTO dishes (name, price, cost) VALUES (?,?,?)",
              (data["name"], data.get("price", 0), data.get("cost", 0)))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/dishes/<int:did>", methods=["DELETE"])
def del_dish(did):
    db = g.db
    db.execute("DELETE FROM dishes WHERE id=?", (did,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/price", methods=["POST"])
def calc_price():
    """按65%毛利计算售价：售价 = 成本 / (1 - 0.65)，向上取整"""
    data = request.get_json(force=True)
    cost = float(data.get("cost", 0))
    margin = float(data.get("margin", 0.65))
    if margin >= 1:
        return jsonify({"ok": False, "msg": "毛利率不能>=100%"}), 400
    price = cost / (1 - margin)
    # 个位向上取整
    price_int = math.ceil(price)
    return jsonify({"ok": True, "cost": cost, "price": price_int, "raw_price": round(price, 2)})


# ===== 全局设置 =====
def get_setting(key, default=None):
    row = g.db.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default


def refresh_package_prices(db):
    """根据 settings.delivery_fee 同步所有 packages.price = base_price + delivery_fee"""
    row = db.execute("SELECT CAST(value AS REAL) as df FROM settings WHERE key='delivery_fee'").fetchone()
    df = row["df"] if row else 0
    db.execute("UPDATE packages SET price = base_price + ?", (df,))


@app.route("/api/settings", methods=["GET"])
def api_get_settings():
    rows = g.db.execute("SELECT * FROM settings ORDER BY key").fetchall()
    data = {r["key"]: {"value": r["value"], "note": r["note"]} for r in rows}
    # 顺便把每个套餐的 base_price / price / delivery 拆分也给前端
    pkgs = g.db.execute("SELECT id, name, base_price, price, price - base_price as delivery FROM packages ORDER BY min_people").fetchall()
    return jsonify({"ok": True, "settings": data, "packages": [dict(r) for r in pkgs]})


@app.route("/api/settings", methods=["POST"])
def api_update_settings():
    """设置多个 key=value，例：{"delivery_fee": "80"}"""
    data = request.get_json(force=True) or {}
    db = g.db
    for k, v in data.items():
        if k == "delivery_fee":
            # 强制数值化，脏输入过滤掉
            try:
                v = str(float(v))
            except (TypeError, ValueError):
                continue
        db.execute("INSERT INTO settings (key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                   (k, v))
    refresh_package_prices(db)
    db.commit()
    return jsonify({"ok": True})


# ===== 历史分析 =====
@app.route("/api/history")
def api_history():
    db = g.db
    # 按月聚合
    rows = db.execute("""
        SELECT strftime('%Y-%m', created_at) as month,
               COUNT(*) as cnt,
               COALESCE(SUM(CASE WHEN status!='cancelled' THEN amount ELSE 0 END),0) as revenue
        FROM orders WHERE deleted_at IS NULL GROUP BY strftime('%Y-%m', created_at)
        ORDER BY month DESC LIMIT 12
    """).fetchall()
    monthly = [dict(r) for r in rows]

    # 按周聚合（近 8 周）
    rows = db.execute("""
        SELECT strftime('%Y-W%W', created_at) as week,
               COUNT(*) as cnt,
               COALESCE(SUM(CASE WHEN status!='cancelled' THEN amount ELSE 0 END),0) as revenue
        FROM orders WHERE deleted_at IS NULL GROUP BY strftime('%Y-W%W', created_at)
        ORDER BY week DESC LIMIT 8
    """).fetchall()
    weekly = list(reversed([dict(r) for r in rows]))

    # 按套餐统计（所有订单中各套餐被点了多少次）
    rows = db.execute("""
        SELECT p.name as pkg_name,
               COUNT(*) as times,
               SUM(op.quantity) as total_qty
        FROM order_packages op
        JOIN orders o ON op.order_id = o.id
        JOIN packages p ON op.package_id = p.id
        WHERE o.status != 'cancelled' AND o.deleted_at IS NULL
        GROUP BY p.id
        ORDER BY total_qty DESC
    """).fetchall()
    packages_pop = [dict(r) for r in rows]

    # 状态分布
    rows = db.execute("""
        SELECT status, COUNT(*) as cnt FROM orders WHERE deleted_at IS NULL GROUP BY status
    """).fetchall()
    status_dist = {r["status"]: r["cnt"] for r in rows}

    # 总览
    rows = db.execute("""
        SELECT COUNT(*) as total,
               COALESCE(SUM(CASE WHEN status!='cancelled' THEN amount ELSE 0 END),0) as total_revenue,
               AVG(CASE WHEN status!='cancelled' THEN amount END) as avg_amount
        FROM orders WHERE deleted_at IS NULL
    """).fetchone()
    overview = dict(rows)

    return jsonify({"ok": True, "monthly": monthly, "weekly": weekly,
                    "packages_pop": packages_pop, "status_dist": status_dist, "overview": overview})


# ===== 财务系统 =====
@app.route("/api/finance/summary")
def finance_summary():
    """财务总览：收入、成本、毛利、押金、应收
    可选参数：date_from, date_to（YYYY-MM-DD），按 booking_date 过滤"""
    db = g.db
    cur = db.cursor()
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    # 构建日期过滤条件（用参数化，带 o. 前缀避免 JOIN 歧义）
    date_clause = "o.status != 'cancelled' AND o.deleted_at IS NULL"
    params = []
    if date_from:
        date_clause += " AND o.booking_date >= ?"
        params.append(date_from)
    if date_to:
        date_clause += " AND o.booking_date <= ?"
        params.append(date_to)
    # 收入（非取消订单）
    r = cur.execute(f"""
        SELECT COUNT(*) as order_cnt,
               COALESCE(SUM(amount),0) as total_revenue,
               COALESCE(SUM(deposit),0) as total_deposit,
               COALESCE(SUM(CASE WHEN payment_status='paid' THEN amount ELSE 0 END),0) as paid_amount,
               COALESCE(SUM(CASE WHEN payment_status='unpaid' THEN amount ELSE 0 END),0) as unpaid_amount,
               COALESCE(SUM(CASE WHEN payment_status='partial' THEN amount ELSE 0 END),0) as partial_amount,
               COALESCE(SUM(CASE WHEN deposit_status='pending' THEN deposit ELSE 0 END),0) as deposit_pending,
               COALESCE(SUM(CASE WHEN deposit_status='returned' THEN deposit ELSE 0 END),0) as deposit_returned,
               COALESCE(SUM(CASE WHEN deposit_status='forfeited' THEN deposit ELSE 0 END),0) as deposit_forfeited
        FROM orders o WHERE {date_clause}
    """, params).fetchone()
    summary = dict(r)
    # 成本计算：每单食材成本 = 各食材用量 × 单价
    r = cur.execute(f"""
        SELECT COALESCE(SUM(pi.per_package * op.quantity * i.cost),0) as food_cost
        FROM order_packages op
        JOIN orders o ON op.order_id = o.id
        JOIN package_ingredients pi ON pi.package_id = op.package_id
        JOIN ingredients i ON pi.ingredient_id = i.id
        WHERE {date_clause}
    """, params).fetchone()
    summary["food_cost"] = r["food_cost"] or 0
    # 工具损耗成本（丢失的工具 × 真实成本）
    r = cur.execute(f"""
        SELECT COALESCE(SUM(tl.lost_qty * t.cost),0) as tool_loss_cost
        FROM tool_loans tl
        JOIN orders o ON tl.order_id = o.id
        JOIN tools t ON tl.tool_id = t.id
        WHERE {date_clause}
    """, params).fetchone()
    summary["tool_loss_cost"] = r["tool_loss_cost"] or 0
    # 配送费收入
    r = cur.execute(f"""
        SELECT COALESCE(SUM(CASE WHEN o.status!='cancelled'
            THEN (o.amount - p.base_price * op.quantity)
            ELSE 0 END),0) as delivery_revenue
        FROM orders o
        JOIN order_packages op ON op.order_id = o.id
        JOIN packages p ON op.package_id = p.id
        WHERE {date_clause}
    """, params).fetchone()
    summary["delivery_revenue"] = r["delivery_revenue"] or 0
    summary["total_cost"] = summary["food_cost"] + summary["tool_loss_cost"]
    summary["gross_profit"] = summary["total_revenue"] - summary["total_cost"]
    summary["margin"] = round(summary["gross_profit"] / max(1, summary["total_revenue"]) * 100, 1)
    # 库存价值（食材+工具）
    r = cur.execute("SELECT COALESCE(SUM(stock * cost),0) as stock_value FROM ingredients").fetchone()
    summary["stock_value"] = r["stock_value"] or 0
    r = cur.execute("SELECT COALESCE(SUM(stock * cost),0) as tool_stock_value FROM tools").fetchone()
    summary["tool_stock_value"] = r["tool_stock_value"] or 0
    return jsonify({"ok": True, "data": summary})


@app.route("/api/finance/orders")
def finance_orders():
    """订单财务明细列表
    可选参数：date_from, date_to（YYYY-MM-DD），按 booking_date 过滤"""
    db = g.db
    cur = db.cursor()
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    date_clause = "1=1"
    params = []
    if date_from:
        date_clause += " AND o.booking_date >= ?"
        params.append(date_from)
    if date_to:
        date_clause += " AND o.booking_date <= ?"
        params.append(date_to)
    rows = cur.execute(f"""
        SELECT o.id, o.booking_date, o.contact_name, o.contact_phone, o.address,
               o.amount, o.deposit, o.payment_status, o.deposit_status, o.status,
               o.created_at,
               (SELECT COALESCE(SUM(pi.per_package * op.quantity * i.cost),0)
                FROM order_packages op
                JOIN package_ingredients pi ON pi.package_id = op.package_id
                JOIN ingredients i ON pi.ingredient_id = i.id
                WHERE op.order_id = o.id) as food_cost,
               (SELECT COALESCE(SUM(tl.lost_qty * t.cost),0)
                FROM tool_loans tl JOIN tools t ON tl.tool_id = t.id
                WHERE tl.order_id = o.id) as tool_loss
        FROM orders o WHERE {date_clause}
        ORDER BY o.booking_date DESC, o.created_at DESC LIMIT 500
    """, params).fetchall()
    orders = []
    for r in rows:
        o = dict(r)
        o["total_cost"] = (o["food_cost"] or 0) + (o["tool_loss"] or 0)
        o["profit"] = (o["amount"] or 0) - o["total_cost"]
        o["margin_pct"] = round(o["profit"] / max(1, o["amount"] or 1) * 100, 1)
        orders.append(o)
    return jsonify({"ok": True, "data": orders})


@app.route("/api/finance/order/<int:oid>/cost-detail")
def finance_order_cost_detail(oid):
    """单个订单的食材成本明细（用于财务页成本编辑）"""
    db = g.db
    cur = db.cursor()
    rows = cur.execute("""
        SELECT pi.per_package, pi.portion_count, pi.cost_only,
               op.quantity as order_qty,
               i.id as ingredient_id, i.name, i.unit, i.cost, i.category
        FROM order_packages op
        JOIN package_ingredients pi ON pi.package_id = op.package_id
        JOIN ingredients i ON pi.ingredient_id = i.id
        WHERE op.order_id = ?
        ORDER BY i.category, i.name
    """, (oid,)).fetchall()
    items = []
    for r in rows:
        total_amount = r["per_package"] * r["order_qty"]
        total_cost = total_amount * (r["cost"] or 0)
        items.append({
            "ingredient_id": r["ingredient_id"],
            "name": r["name"], "unit": r["unit"],
            "per_package": r["per_package"],
            "portion_count": r["portion_count"],
            "order_qty": r["order_qty"],
            "cost_only": r["cost_only"],
            "total_amount": round(total_amount, 2),
            "unit_cost": r["cost"] or 0,
            "total_cost": round(total_cost, 2),
            "category": r["category"],
        })
    return jsonify({"ok": True, "data": items})


@app.route("/api/finance/payment/<int:oid>", methods=["POST"])
def update_payment(oid):
    """更新订单货款状态"""
    data = request.get_json(force=True)
    status = data.get("payment_status", "unpaid")
    db = g.db
    db.execute("UPDATE orders SET payment_status=? WHERE id=?", (status, oid))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/finance/deposit/<int:oid>", methods=["POST"])
def update_deposit(oid):
    """更新押金状态"""
    data = request.get_json(force=True)
    status = data.get("deposit_status", "pending")
    db = g.db
    db.execute("UPDATE orders SET deposit_status=? WHERE id=?", (status, oid))
    db.commit()
    return jsonify({"ok": True})


# ===== 打印备餐单（按食材汇总）=====
@app.route("/api/print/menu")
def print_menu():
    """返回待备/备餐中订单的食材+工具汇总，适合打印"""
    import json as _json
    db = g.db
    cur = db.cursor()
    cur.execute("SELECT id, booking_date, booking_time, address, contact_name, status "
                "FROM orders WHERE status IN ('pending','preparing') AND deleted_at IS NULL ORDER BY id")
    orders = [dict(r) for r in cur.fetchall()]

    order_details = []
    for o in orders:
        cur.execute("""
            SELECT op.quantity, p.name, p.id
            FROM order_packages op
            JOIN packages p ON op.package_id = p.id
            WHERE op.order_id = ?
        """, (o["id"],))
        pkgs = [dict(r) for r in cur.fetchall()]
        order_details.append({**o, "packages": pkgs})

    # 直接调用 calc_merged_prep，与备餐页数据源一致
    from calculator import calc_merged_prep, CATEGORY_ORDER, CATEGORY_LABEL, \
        PACKAGING_CATEGORY, UTENSIL_CATEGORY, SAUCE_LIKE_CATS, \
        _subcategorize_meat, _get_fixed_sort_index, sort_tools
    merged = calc_merged_prep([o["id"] for o in orders])
    # merged["ingredients"] 是按分类分组的 dict: {cat: [item, ...]}
    ing_by_cat = merged["ingredients"]
    tools_sorted = merged["tools"]

    # 渲染成一个可打印的 HTML（独立页面，无 nav，适合 A4 打印）
    import datetime
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    order_lines = ""
    for o in order_details:
        pkg_str = ", ".join(f"{p['name']}×{p['quantity']}" for p in o["packages"])
        order_lines += f"""
        <tr><td>{o['id']}</td><td>{o['booking_date']} {o['booking_time'] or ''}</td>
            <td>{o['address'] or ''}</td><td>{o['contact_name'] or ''}</td>
            <td>{pkg_str}</td></tr>"""

    # 分类配色（与备餐页/菜单页一致）
    cat_colors = {
        "beef": "#8e1e1a", "pork": "#c75d3e", "chicken": "#c9962b",
        "vegetable": "#3a8a3a", "side": "#5a8a3a", "sauce": "#7a5a3a",
        "drink": "#666666", "staple": "#888888", "other": "#6b4f3a",
        "packaging": "#d35400", "utensil": "#8e44ad", "tool": "#4a4a4a",
    }
    # 食材按分类渲染（与备餐/菜单页顺序一致：含packaging/utensil，不再单独处理）
    ing_rows = ""
    for cat in CATEGORY_ORDER:
        items = ing_by_cat.get(cat, [])
        if not items:
            continue
        label = CATEGORY_LABEL.get(cat, cat)
        color = cat_colors.get(cat, "#6b4f3a")
        ing_rows += f'<tr class="cat-row"><td colspan="5" style="background:{color};color:#fff;font-weight:700;padding:7px 10px;font-size:14px;">{label} · {len(items)}项</td></tr>'
        for i in items:
            tot_portions = i.get("total_portions", 0)
            per_size = i.get("portion_size", 0)
            unit = i.get("unit", "") or ""
            # packaging/utensil/tool(无克数列) 与菜单页一致：克数显示"—"
            is_simple = cat in (PACKAGING_CATEGORY, UTENSIL_CATEGORY)
            if is_simple:
                tot_portions = i.get("total_packages", 0)
                per_size = None  # 不显示克数
            total_val = round(i["total"], 2) if i.get("total") else 0
            total_str = f"{total_val}{unit}" if total_val else "-"
            ing_rows += (
                f'<tr><td class="name-cell">{i["name"]}</td>'
                f'<td class="num-cell">{round(tot_portions) if tot_portions else "-"}</td>'
                f'<td class="num-cell">{round(per_size,2) if per_size else "—"}</td>'
                f'<td class="num-cell total-cell">{total_str}</td>'
                f'<td>{unit}</td></tr>'
            )

    # 工具单独渲染（与备餐页一致：🔧 工具分类，克数列显示"—"，单位默认"个"）
    tool_rows = ""
    if tools_sorted:
        tool_rows += f'<tr class="cat-row"><td colspan="5" style="background:#4a4a4a;color:#fff;font-weight:700;padding:7px 10px;font-size:14px;">🔧 工具 · {len(tools_sorted)}项</td></tr>'
        for t in tools_sorted:
            t_total = round(t.get("total", 0))
            tool_rows += (
                f'<tr><td class="name-cell">{t["name"]}</td>'
                f'<td class="num-cell">{t_total}</td>'
                f'<td class="num-cell">—</td>'
                f'<td class="num-cell total-cell">{t_total}个</td>'
                f'<td>个</td></tr>'
            )

    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<title>备餐单 · {now}</title>
<style>
  @media print{{ @page{{ size:A4; margin:12mm }} }}
  body{{ font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif; color:#2b1e14; padding:20px; max-width:800px; margin:0 auto; }}
  h1{{ font-size:22px; margin:0 0 4px; }}
  .meta{{ color:#8a7f75; font-size:13px; margin-bottom:14px; }}
  .sec-title{{ font-size:15px; font-weight:700; margin:18px 0 8px; padding-bottom:4px; border-bottom:2px solid #c0392b; color:#c0392b; }}
  table{{ width:100%; border-collapse:collapse; font-size:13px; }}
  thead th{{ background:#f3efe8; text-align:left; padding:7px 10px; font-weight:700; color:#666; }}
  thead th.num-col{{ text-align:center; }}
  tbody td{{ padding:8px 10px; border-bottom:1px solid #efe8dd; }}
  tbody td.name-cell{{ font-weight:700; }}
  tbody td.num-cell{{ text-align:center; color:#444; font-weight:600; }}
  tbody td.total-cell{{ text-align:right; color:#c0392b; font-weight:800; font-family:"DIN","Helvetica Neue",sans-serif; }}
  tbody tr.cat-row td{{ padding:7px 10px; }}
  .print-btn{{ margin:12px 6px 12px 0; padding:10px 24px; background:#c0392b; color:#fff; border:none; border-radius:6px; font-size:15px; cursor:pointer; font-weight:700; }}
  .export-btn{{ margin:12px 6px 12px 0; padding:10px 24px; background:#27ae60; color:#fff; border:none; border-radius:6px; font-size:15px; cursor:pointer; font-weight:700; }}
  .back-btn{{ margin:12px 0; padding:10px 24px; background:#fff; color:#2b1e14; border:1.5px solid #c0392b; border-radius:6px; font-size:15px; cursor:pointer; text-decoration:none; display:inline-block; }}
  .back-btn:hover{{ background:#fdecea; }}
  @media print{{ .print-btn, .export-btn, .back-btn{{ display:none }} }}
</style>
</head>
<body>
<a class="back-btn" href="javascript:history.length>1?history.back():'/'">← 返回</a>
<button class="print-btn" onclick="window.print()">🖨️ 打印 / 存为PDF</button>
<button class="export-btn" onclick="exportCSV()">📥 导出Excel(CSV)</button>
<h1>🔥 刘和牛户外烤肉 · 备餐单</h1>
<div class="meta">生成时间：{now} · 共 {len(orders)} 单待备 / 备餐中</div>

<div class="sec-title">📋 订单清单</div>
<table><thead><tr><th>#</th><th>时间</th><th>地址</th><th>联系人</th><th>套餐</th></tr></thead>
<tbody>{order_lines or '<tr><td colspan=5 style="text-align:center;color:#aaa;padding:20px">暂无待备订单</td></tr>'}</tbody></table>

<div class="sec-title">🥩 备餐核对清单</div>
<table><thead><tr>
  <th>名字</th><th class="num-col">份数</th><th class="num-col">克数</th><th class="num-col">总量</th><th>单位</th>
</tr></thead>
<tbody>{ing_rows or '<tr><td colspan=5 style="text-align:center;color:#aaa;padding:20px">—</td></tr>'}{tool_rows}</tbody></table>
<script>
function exportCSV(){{
  var rows=[["分类","食材","份数","克数","总量","单位"]];
  document.querySelectorAll("table tbody tr").forEach(function(tr){{
    if(tr.classList.contains("cat-row")) return;
    var cells=tr.querySelectorAll("td");
    if(cells.length<5) return;
    rows.push([tr.parentElement.previousElementSibling? "":"",
      cells[0]?.textContent||"", cells[1]?.textContent||"",
      cells[2]?.textContent||"", cells[3]?.textContent||"",
      cells[4]?.textContent||""]);
  }});
  var csv="\\uFEFF";
  rows.forEach(function(r){{ csv+=r.map(function(c){{ return '"'+(c||"").replace(/"/g,'""')+'"'; }}).join(",")+";"; }});
  var blob=new Blob([csv],{{type:"text/csv;charset=utf-8;"}});
  var a=document.createElement("a");
  a.href=URL.createObjectURL(blob);
  a.download="备餐单_"+new Date().toISOString().slice(0,10)+".csv";
  a.click();
}}
</script>
</body></html>"""
    return html


# ===== 工具借出归还 =====
@app.route("/api/orders/<int:oid>/loans")
def get_order_loans(oid):
    """获取订单的工具借出记录"""
    db = g.db
    cur = db.cursor()
    cur.execute("""
        SELECT tl.*, t.name as tool_name
        FROM tool_loans tl JOIN tools t ON tl.tool_id = t.id
        WHERE tl.order_id = ? ORDER BY tl.id
    """, (oid,))
    loans = [dict(r) for r in cur.fetchall()]
    return jsonify({"ok": True, "data": loans})


@app.route("/api/orders/<int:oid>/tools/return", methods=["POST"])
def return_tools(oid):
    """归还工具。body: {items: [{loan_id, returned_qty, lost_qty, note}]}
    lost_qty > 0 时自动从工具总库存扣除（报损）。
    """
    data = request.get_json(force=True)
    items = data.get("items", [])
    db = g.db
    cur = db.cursor()
    done = 0
    for it in items:
        lid = it.get("loan_id")
        returned = float(it.get("returned_qty", 0))
        lost = float(it.get("lost_qty", 0))
        note = it.get("note", "")
        cur.execute("SELECT * FROM tool_loans WHERE id=? AND order_id=?", (lid, oid))
        loan = cur.fetchone()
        if not loan:
            continue
        total = loan["quantity"]
        new_returned = min(total, returned)
        new_lost = min(total - new_returned, lost)
        if new_returned + new_lost >= total:
            status = "returned" if new_lost == 0 else "lost"
        else:
            status = "partial"
        cur.execute("""
            UPDATE tool_loans SET returned_qty=?, lost_qty=?, status=?, note=?,
            returned_at=datetime('now','localtime') WHERE id=?
        """, (new_returned, new_lost, status, note, lid))
        # 丢失工具从总库存扣除
        if new_lost > 0:
            cur.execute("UPDATE tools SET stock = stock - ? WHERE id=?", (new_lost, loan["tool_id"]))
            cur.execute("""
                INSERT INTO stock_logs (item_type, item_id, delta, reason)
                VALUES ('tool',?,?,?)
            """, (loan["tool_id"], -new_lost, f"订单#{oid} 工具丢失/损坏"))
        done += 1
    db.commit()
    return jsonify({"ok": True, "done": done})


# ===== 备餐合并视图 =====
@app.route("/api/prep/merged")
def prep_merged():
    """合并备餐视图。默认取当天 pending/preparing 订单；可传 ?date=YYYY-MM-DD 或 ?ids=1,2,3"""
    import datetime
    db = g.db
    cur = db.cursor()
    ids_param = request.args.get("ids")
    if ids_param:
        order_ids = [int(x) for x in ids_param.split(",") if x.strip().isdigit()]
    else:
        from calculator import today_cst as _today_cst
        date = request.args.get("date") or _today_cst().isoformat()
        cur.execute("""
            SELECT id FROM orders
            WHERE status IN ('pending','preparing') AND deleted_at IS NULL AND booking_date = ?
            ORDER BY booking_time
        """, (date,))
        order_ids = [r["id"] for r in cur.fetchall()]
    data = calc_merged_prep(order_ids)
    return jsonify({"ok": True, "data": data, "order_ids": order_ids})


@app.route("/api/prep/dates")
def prep_dates():
    """列出所有有 pending/preparing 订单的日期，供前端做日期快捷切换"""
    from calculator import today_cst as _today_cst
    db = g.db
    cur = db.cursor()
    cur.execute("""
        SELECT booking_date, COUNT(*) as cnt
        FROM orders
        WHERE status IN ('pending','preparing') AND deleted_at IS NULL AND booking_date IS NOT NULL AND booking_date != ''
        GROUP BY booking_date
        ORDER BY booking_date
    """)
    today = _today_cst().isoformat()
    dates = []
    for r in cur.fetchall():
        d = dict(r)
        d["is_today"] = (r["booking_date"] == today)
        d["is_past"] = (r["booking_date"] < today)
        dates.append(d)
    return jsonify({"ok": True, "dates": dates, "today": today})


# ===== 备餐勾选 =====
@app.route("/api/prep/check", methods=["POST"])
def prep_check():
    """勾选/取消备餐项。
    body: {order_id or order_ids, item_type, item_id, checked}
    order_ids: 数组，用于合并视图一次勾选多个订单的同一项。
    """
    data = request.get_json(force=True)
    itype = data.get("item_type")
    iid = data.get("item_id")
    checked = 1 if data.get("checked") else 0
    oids = data.get("order_ids")
    if oids is None:
        oids = [data.get("order_id")]
    oids = [int(x) for x in oids if x]
    if not oids or itype not in ("ingredient", "tool") or not iid:
        return jsonify({"ok": False, "msg": "参数错误"}), 400
    db = g.db
    cur = db.cursor()
    import datetime as _dt
    ts = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for oid in oids:
        cur.execute("""
            SELECT id, quantity FROM prep_checklist
            WHERE order_id=? AND item_type=? AND item_id=?
        """, (oid, itype, iid))
        row = cur.fetchone()
        if row:
            cur.execute("UPDATE prep_checklist SET checked=?, checked_at=? WHERE id=?",
                        (checked, ts if checked else None, row["id"]))
        else:
            qty = 0
            if itype == "ingredient":
                cur.execute("SELECT per_package FROM package_ingredients WHERE ingredient_id=? LIMIT 1", (iid,))
                r = cur.fetchone()
                qty = r["per_package"] if r else 0
            else:
                cur.execute("SELECT per_package FROM package_tools WHERE tool_id=? LIMIT 1", (iid,))
                r = cur.fetchone()
                qty = r["per_package"] if r else 0
            cur.execute("""
                INSERT INTO prep_checklist (order_id, item_type, item_id, quantity, checked, checked_at)
                VALUES (?,?,?,?,?,?)
            """, (oid, itype, iid, qty, checked, ts if checked else None))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/prep/check/batch", methods=["POST"])
def prep_check_batch():
    """批量勾选：给指定订单的所有备餐项设置 checked 状态
    body: {order_ids: [...], checked: 0/1}
    """
    data = request.get_json(force=True)
    order_ids = data.get("order_ids", [])
    checked = 1 if data.get("checked") else 0
    db = g.db
    cur = db.cursor()
    import datetime as _dt
    ts = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for oid in order_ids:
        # 确保所有需求项都有记录
        req = calc_order_requirements(oid)
        for ing in req["ingredients"]:
            cur.execute("SELECT id FROM prep_checklist WHERE order_id=? AND item_type='ingredient' AND item_id=?",
                        (oid, ing["id"]))
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO prep_checklist (order_id, item_type, item_id, quantity, checked)
                    VALUES (?,?,?,?,?)
                """, (oid, "ingredient", ing["id"], ing["need"], 0))
        for tl in req["tools"]:
            cur.execute("SELECT id FROM prep_checklist WHERE order_id=? AND item_type='tool' AND item_id=?",
                        (oid, tl["id"]))
            if not cur.fetchone():
                cur.execute("""
                    INSERT INTO prep_checklist (order_id, item_type, item_id, quantity, checked)
                    VALUES (?,?,?,?,?)
                """, (oid, "tool", tl["id"], tl["need"], 0))
        # 批量更新该订单所有项
        cur.execute("UPDATE prep_checklist SET checked=?, checked_at=? WHERE order_id=?",
                    (checked, ts if checked else None, oid))
    db.commit()
    return jsonify({"ok": True})


if __name__ == "__main__":
    import os
    port = int(os.environ.get("PORT", 8000))
    app.run(host="0.0.0.0", port=port, debug=False)
